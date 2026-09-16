"""Coverage for `relay_core.rag.parsers` (docs/system-design.md section 11.1): each format
parser turns raw bytes into the same flat `list[Block]` shape the chunker expects, and the PDF
parser additionally flags scanned (image-only) pages via `needs_ocr` for the ingestion
pipeline's Gemini-fallback branch.
"""

import io

import pymupdf
import pytest
from docx import Document as DocxDocument

from relay_core.rag.parsers import parse_document, parse_docx, parse_markdown, parse_pdf, parse_text


def test_parse_text_splits_on_blank_lines() -> None:
    raw = b"First paragraph.\n\nSecond paragraph.\n\n\nThird paragraph."
    doc = parse_text(raw)
    assert [b.text for b in doc.blocks] == [
        "First paragraph.",
        "Second paragraph.",
        "Third paragraph.",
    ]
    assert all(b.kind == "paragraph" for b in doc.blocks)


def test_parse_markdown_extracts_headings_with_levels_and_paragraphs() -> None:
    raw = b"# Title\n\nIntro text.\n\n## Sub\n\nMore text.\n"
    doc = parse_markdown(raw)
    kinds = [(b.kind, b.level, b.text) for b in doc.blocks]
    assert kinds == [
        ("heading", 1, "Title"),
        ("paragraph", None, "Intro text."),
        ("heading", 2, "Sub"),
        ("paragraph", None, "More text."),
    ]


def test_parse_markdown_keeps_fenced_code_as_a_paragraph_block() -> None:
    raw = b"# Title\n\n```\nprint('hi')\n```\n"
    doc = parse_markdown(raw)
    assert any(b.kind == "paragraph" and "print" in b.text for b in doc.blocks)


def test_parse_docx_detects_heading_styles_and_tables_in_document_order() -> None:
    docx_doc = DocxDocument()
    docx_doc.add_heading("Policy", level=1)
    docx_doc.add_paragraph("Body text.")
    docx_doc.add_heading("Exceptions", level=2)
    table = docx_doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Required"
    table.cell(1, 0).text = "Admin"
    table.cell(1, 1).text = "Yes"
    buf = io.BytesIO()
    docx_doc.save(buf)

    parsed = parse_docx(buf.getvalue())
    kinds = [(b.kind, b.level) for b in parsed.blocks]
    assert kinds == [
        ("heading", 1),
        ("paragraph", None),
        ("heading", 2),
        ("table", None),
    ]
    table_block = next(b for b in parsed.blocks if b.kind == "table")
    assert "| Role | Required |" in table_block.text
    assert "| Admin | Yes |" in table_block.text


def test_parse_docx_skips_empty_paragraphs() -> None:
    docx_doc = DocxDocument()
    docx_doc.add_paragraph("")
    docx_doc.add_paragraph("Real content.")
    buf = io.BytesIO()
    docx_doc.save(buf)

    parsed = parse_docx(buf.getvalue())
    assert [b.text for b in parsed.blocks] == ["Real content."]


def test_parse_pdf_extracts_text_and_reports_the_page_count() -> None:
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page()
    page.insert_text((72, 72), "Policy Title", fontsize=20)
    page.insert_text((72, 110), "Body text goes here in the document.", fontsize=11)
    raw = pdf_doc.tobytes()

    parsed = parse_pdf(raw)
    assert parsed.page_count == 1
    assert parsed.needs_ocr is False
    assert any(b.kind == "heading" for b in parsed.blocks)
    assert any(b.kind == "paragraph" and "Body text" in b.text for b in parsed.blocks)
    assert all(b.page == 1 for b in parsed.blocks)


def test_parse_pdf_flags_a_blank_page_as_needing_ocr() -> None:
    pdf_doc = pymupdf.open()
    pdf_doc.new_page()
    parsed = parse_pdf(pdf_doc.tobytes())
    assert parsed.needs_ocr is True
    assert parsed.blocks == []


def test_parse_document_dispatches_on_mime_type() -> None:
    assert parse_document("text/plain", b"Hello.").blocks[0].text == "Hello."
    assert parse_document("text/markdown", b"# T\n\nBody.\n").blocks[0].kind == "heading"


def test_parse_document_rejects_an_unsupported_mime_type() -> None:
    with pytest.raises(ValueError, match="Unsupported mime type"):
        parse_document("application/zip", b"whatever")
